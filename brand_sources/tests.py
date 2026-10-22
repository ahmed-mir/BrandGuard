import hashlib
import io
import secrets
import zipfile
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.core.files.storage import storages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError

from workspaces.models import Workspace, WorkspaceMembership

from .demo import LUMADESK_SOURCE_PACK
from .ingestion import DocumentImportError, extract_uploaded_document
from .models import ApprovedFact, BrandDocumentImport, BrandRule, BrandSource, SourceVersion, policy_snapshot_hash
from .services import (
    DocumentImportPersistenceError,
    active_policy_versions,
    approve_source_version,
    confirm_imported_text,
    create_draft_version,
    create_source,
    import_source_document,
    policy_snapshot,
)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class BrandSourceWorkflowTests(TestCase):
    def setUp(self):
        self.workspace = Workspace.objects.create(name="LumaDesk", slug="brand-source-test")
        self.manager = self._member("policy_manager", WorkspaceMembership.Role.BRAND_MANAGER)
        self.admin = self._member("policy_admin", WorkspaceMembership.Role.ADMIN)
        self.creator = self._member("policy_creator", WorkspaceMembership.Role.CREATOR)
        self.reviewer = self._member("policy_reviewer", WorkspaceMembership.Role.REVIEWER)
        self.source, self.version = create_source(
            workspace=self.workspace,
            owner=self.manager,
            title="Naming rules",
            source_type=BrandSource.SourceType.GUIDELINES,
            description="Approved naming reference.",
            source_text="Official product name: LumaDesk. Do not write Luma Desk as two words.",
        )

    def _member(self, username, role, workspace=None):
        user = get_user_model().objects.create_user(
            username=username,
            email=f"{username}@example.test",
            password=secrets.token_urlsafe(24),
        )
        WorkspaceMembership.objects.create(user=user, workspace=workspace or self.workspace, role=role)
        return user

    def _create_rule(self, version=None, **overrides):
        version = version or self.version
        values = {
            "source_version": version,
            "identifier": "TERM-001",
            "category": BrandRule.Category.TERMINOLOGY,
            "severity": BrandRule.Severity.MEDIUM,
            "description": "Use the official product name.",
            "rule_type": BrandRule.RuleType.PROHIBITED_TERM,
            "matcher_data": {"type": "PROHIBITED_TERM", "term": "Luma Desk", "preferred": "LumaDesk", "case_sensitive": False, "scope": "ALL"},
            "evidence_excerpt": "Do not write Luma Desk as two words.",
        }
        values.update(overrides)
        return BrandRule.objects.create(**values)

    def _create_fact(self, version=None, **overrides):
        version = version or self.version
        values = {
            "source_version": version,
            "identifier": "FACT-02",
            "subject": "Task completion speed",
            "statement": "Up to 20% faster task completion in our internal pilot; results vary.",
            "comparison_operator": ApprovedFact.ComparisonOperator.UP_TO,
            "comparison_value": "20",
            "comparison_unit": "%",
            "evidence_excerpt": "Up to 20% faster task completion in our internal pilot; results vary.",
        }
        values.update(overrides)
        return ApprovedFact.objects.create(**values)

    def test_source_creation_stores_workspace_owner_and_first_draft(self):
        self.client.force_login(self.manager)
        response = self.client.post(reverse("app:brand_sources:create"), {
            "title": "Channel disclaimer",
            "source_type": BrandSource.SourceType.POLICY,
            "description": "Channel guidance.",
            "source_text": "Include Results vary. with pilot claims.",
            "effective_date": "",
        })
        self.assertEqual(response.status_code, 302, response.content.decode("utf-8"))
        source = BrandSource.objects.get(title="Channel disclaimer")
        version = source.versions.get(version_number=1)
        self.assertEqual(source.workspace, self.workspace)
        self.assertEqual(source.owner, self.manager)
        self.assertEqual(version.status, SourceVersion.Status.DRAFT)
        self.assertEqual(version.content_hash, hashlib.sha256(version.source_text.encode()).hexdigest())

    def test_source_library_detail_and_editor_pages_render_with_named_routes(self):
        self._create_rule()
        approve_source_version(source_id=self.source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        self.client.force_login(self.manager)
        routes = [
            reverse("app:brand_sources:list"),
            reverse("app:brand_sources:detail", args=[self.source.pk]),
            reverse("app:brand_sources:version_detail", args=[self.source.pk, 1]),
            reverse("app:brand_sources:version_create", args=[self.source.pk]),
        ]
        for route in routes:
            with self.subTest(route=route):
                self.assertIn(self.client.get(route).status_code, {200, 302})
        self.client.post(reverse("app:brand_sources:version_create", args=[self.source.pk]), {
            "source_text": self.version.source_text + " Current guidance.", "effective_date": "",
        })
        draft_routes = [
            reverse("app:brand_sources:version_edit", args=[self.source.pk, 2]),
            reverse("app:brand_sources:rule_create", args=[self.source.pk, 2]),
            reverse("app:brand_sources:fact_create", args=[self.source.pk, 2]),
        ]
        for route in draft_routes:
            with self.subTest(route=route):
                self.assertEqual(self.client.get(route).status_code, 200)

    def test_invalid_source_text_is_rejected(self):
        self.client.force_login(self.manager)
        response = self.client.post(reverse("app:brand_sources:create"), {
            "title": "Empty", "source_type": "POLICY", "description": "", "source_text": "   ", "effective_date": "",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Add the full authoritative source text")
        self.assertFalse(BrandSource.objects.filter(title="Empty").exists())

    def test_rule_form_creates_typed_rule_with_exact_source_evidence(self):
        self.client.force_login(self.manager)
        response = self.client.post(reverse("app:brand_sources:rule_create", args=[self.source.pk, 1]), {
            "identifier": "TERM-003",
            "category": BrandRule.Category.TERMINOLOGY,
            "severity": BrandRule.Severity.HIGH,
            "rule_type": BrandRule.RuleType.PROHIBITED_TERM,
            "description": "Use one-word product spelling.",
            "term": "Luma Desk",
            "preferred_term": "LumaDesk",
            "case_sensitive": "",
            "scope": "ALL",
            "evidence_excerpt": "Do not write Luma Desk as two words.",
            "evidence_start": "",
            "evidence_end": "",
        })
        self.assertEqual(response.status_code, 302, response.content.decode("utf-8"))
        rule = BrandRule.objects.get(identifier="TERM-003")
        self.assertEqual(rule.matcher_data["term"], "Luma Desk")
        self.assertEqual(rule.matcher_data["preferred"], "LumaDesk")
        self.assertEqual(rule.source_version, self.version)

    def test_rule_rejects_missing_evidence_and_malformed_or_executable_matchers(self):
        self.client.force_login(self.manager)
        response = self.client.post(reverse("app:brand_sources:rule_create", args=[self.source.pk, 1]), {
            "identifier": "TERM-004", "category": "TERMINOLOGY", "severity": "MEDIUM",
            "rule_type": "PROHIBITED_TERM", "description": "Bad evidence", "term": "X",
            "scope": "ALL", "evidence_excerpt": "not in the source", "evidence_start": "", "evidence_end": "",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "does not appear in this draft")
        with self.assertRaises(ValidationError):
            self._create_rule(identifier="TERM-005", matcher_data={"type": "PROHIBITED_TERM", "expression": "__import__('os')"})
        with self.assertRaises(ValidationError):
            self._create_rule(identifier="TERM-006", matcher_data={"type": "PROHIBITED_TERM", "term": "X", "scope": "ALL", "regex": ".*"})

    def test_evidence_offsets_must_match_the_exact_excerpt(self):
        start = self.version.source_text.index("LumaDesk")
        end = start + len("LumaDesk")
        rule = self._create_rule(
            identifier="TERM-007",
            evidence_excerpt="LumaDesk",
            evidence_start=start,
            evidence_end=end,
        )
        self.assertEqual(self.version.source_text[rule.evidence_start:rule.evidence_end], rule.evidence_excerpt)
        with self.assertRaises(ValidationError):
            self._create_rule(identifier="TERM-008", evidence_excerpt="LumaDesk", evidence_start=start, evidence_end=end - 1)

    def test_fact_form_saves_normalized_fact_and_linked_approved_fact_rule(self):
        source, version = create_source(
            workspace=self.workspace, owner=self.manager, title="Facts", source_type="FACT_SHEET",
            description="", source_text="Up to 20% faster task completion in our internal pilot; results vary.",
        )
        self.client.force_login(self.manager)
        response = self.client.post(reverse("app:brand_sources:fact_create", args=[source.pk, 1]), {
            "identifier": "FACT-02", "subject": "Task completion speed",
            "statement": "Up to 20% faster task completion in our internal pilot; results vary.",
            "comparison_operator": "UP_TO", "comparison_value": "20", "comparison_max": "",
            "comparison_unit": "%", "evidence_excerpt": "Up to 20% faster task completion in our internal pilot; results vary.",
            "evidence_start": "", "evidence_end": "", "effective_date": "",
        })
        self.assertEqual(response.status_code, 302, response.content.decode("utf-8"))
        fact = version.facts.get(identifier="FACT-02")
        rule = version.rules.get(identifier="FACT-02")
        self.assertEqual(fact.comparison_value, 20)
        self.assertEqual(rule.rule_type, BrandRule.RuleType.APPROVED_FACT)
        self.assertEqual(rule.approved_fact, fact)
        self.assertEqual(rule.matcher_data, {"type": "APPROVED_FACT", "fact_identifier": "FACT-02"})

    def test_fact_rejects_evidence_that_is_not_verbatim_source_text(self):
        self.client.force_login(self.manager)
        response = self.client.post(reverse("app:brand_sources:fact_create", args=[self.source.pk, 1]), {
            "identifier": "FACT-02", "subject": "Task completion", "statement": "A claim.",
            "comparison_operator": "", "comparison_value": "", "comparison_max": "", "comparison_unit": "",
            "evidence_excerpt": "A citation that was not sourced", "evidence_start": "", "evidence_end": "", "effective_date": "",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "does not appear in this draft")
        self.assertFalse(self.version.facts.exists())

    def test_new_draft_version_clones_structured_policy_and_preserves_old_version(self):
        rule = self._create_rule()
        fact_source_text = self.version.source_text + "\nUp to 20% faster task completion in our internal pilot; results vary."
        self.version.source_text = fact_source_text
        self.version.save()
        fact = self._create_fact(evidence_excerpt="Up to 20% faster task completion in our internal pilot; results vary.")
        fact_rule = BrandRule.objects.create(
            source_version=self.version, approved_fact=fact, identifier=fact.identifier, category="FACT", severity="MEDIUM",
            description="Approved factual statement about task completion.", rule_type="APPROVED_FACT",
            matcher_data={"type": "APPROVED_FACT", "fact_identifier": fact.identifier}, evidence_excerpt=fact.evidence_excerpt,
        )
        approve_source_version(source_id=self.source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        self.client.force_login(self.manager)
        response = self.client.post(reverse("app:brand_sources:version_create", args=[self.source.pk]), {
            "source_text": self.version.source_text + " New guidance.", "effective_date": "",
        })
        self.assertEqual(response.status_code, 302, response.content.decode("utf-8"))
        draft = self.source.versions.get(version_number=2)
        self.assertEqual(draft.status, SourceVersion.Status.DRAFT)
        self.assertEqual(draft.rules.count(), 2)
        self.assertEqual(draft.facts.count(), 1)
        self.assertNotEqual(draft.rules.get(identifier=rule.identifier).pk, rule.pk)
        self.assertNotEqual(draft.facts.get(identifier=fact.identifier).pk, fact.pk)
        self.assertEqual(draft.rules.get(identifier=fact_rule.identifier).approved_fact.source_version, draft)
        self.version.refresh_from_db()
        self.assertEqual(self.version.status, SourceVersion.Status.APPROVED)
        self.assertEqual(self.version.approved_by, self.manager)

    def test_approved_version_is_immutable_and_records_actor_and_timestamp(self):
        self._create_rule()
        approved = approve_source_version(source_id=self.source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        self.assertEqual(approved.status, SourceVersion.Status.APPROVED)
        self.assertEqual(approved.approved_by, self.manager)
        self.assertIsNotNone(approved.approved_at)
        self.workspace.refresh_from_db()
        self.assertEqual(self.workspace.policy_revision, 1)
        self.assertEqual(self.workspace.policy_snapshot_hash, policy_snapshot_hash(self.workspace))
        self.assertEqual(active_policy_versions(self.workspace).get(), approved)
        snapshot = policy_snapshot(self.workspace)
        self.assertEqual(snapshot["revision"], 1)
        self.assertEqual(snapshot["versions"][0]["version_id"], str(approved.pk))

        approved.source_text = "Unauthorized rewrite"
        with self.assertRaises(ValidationError):
            approved.save()
        rule = approved.rules.get(identifier="TERM-001")
        rule.description = "Unauthorized edit"
        with self.assertRaises(ValidationError):
            rule.save()
        with self.assertRaises(ValidationError):
            rule.delete()
        with self.assertRaises(ValidationError):
            approved.delete()
        with self.assertRaises(ValidationError):
            self.source.title = "Rewritten source"
            self.source.save()

    def test_new_approval_atomically_supersedes_previous_and_changes_policy_revision(self):
        self._create_rule()
        first = approve_source_version(source_id=self.source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        prior_hash = self.workspace.policy_snapshot_hash
        draft, created = create_draft_version(source=self.source, workspace=self.workspace)
        self.assertTrue(created)
        draft.source_text += " A later approved sentence."
        draft.save()
        second = approve_source_version(source_id=self.source.pk, version_number=2, workspace=self.workspace, actor=self.admin)
        first.refresh_from_db()
        self.workspace.refresh_from_db()
        self.assertEqual(first.status, SourceVersion.Status.SUPERSEDED)
        self.assertEqual(first.approved_by, self.manager)
        self.assertIsNotNone(first.approved_at)
        self.assertEqual(second.status, SourceVersion.Status.APPROVED)
        self.assertEqual(self.workspace.policy_revision, 2)
        self.assertNotEqual(self.workspace.policy_snapshot_hash, prior_hash)
        self.assertEqual(list(active_policy_versions(self.workspace)), [second])
        self.assertEqual(self.source.versions.count(), 2)

    def test_invalid_persisted_evidence_prevents_approval_without_partial_supersession(self):
        self._create_rule()
        first = approve_source_version(source_id=self.source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        draft, _ = create_draft_version(source=self.source, workspace=self.workspace)
        bad_rule = draft.rules.get(identifier="TERM-001")
        BrandRule.objects.filter(pk=bad_rule.pk).update(evidence_excerpt="not source-backed")
        self.client.force_login(self.admin)
        response = self.client.post(reverse("app:brand_sources:approve", args=[self.source.pk, 2]))
        self.assertEqual(response.status_code, 302)
        self.assertIn("does not appear in the selected source version", str(list(response.wsgi_request._messages)))
        draft.refresh_from_db()
        first.refresh_from_db()
        self.assertEqual(draft.status, SourceVersion.Status.DRAFT)
        self.assertEqual(first.status, SourceVersion.Status.APPROVED)
        self.workspace.refresh_from_db()
        self.assertEqual(self.workspace.policy_revision, 1)

    def test_source_editor_and_rule_mutations_are_manager_or_admin_only(self):
        for user in (self.creator, self.reviewer):
            with self.subTest(role=user.username):
                self.client.force_login(user)
                self.assertEqual(self.client.get(reverse("app:brand_sources:create")).status_code, 403)
                self.assertEqual(self.client.get(reverse("app:brand_sources:version_edit", args=[self.source.pk, 1])).status_code, 403)
        for user in (self.manager, self.admin):
            with self.subTest(role=user.username):
                self.client.force_login(user)
                self.assertEqual(self.client.get(reverse("app:brand_sources:create")).status_code, 200)
                self.assertEqual(self.client.get(reverse("app:brand_sources:version_edit", args=[self.source.pk, 1])).status_code, 200)

    def test_creators_and_reviewers_can_read_approved_sources_but_not_drafts(self):
        self._create_rule()
        approve_source_version(source_id=self.source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        draft, _ = create_draft_version(source=self.source, workspace=self.workspace)
        for user in (self.creator, self.reviewer):
            self.client.force_login(user)
            listing = self.client.get(reverse("app:brand_sources:list"))
            self.assertEqual(listing.status_code, 200)
            self.assertContains(listing, "Naming rules")
            self.assertNotContains(listing, "Draft · v2")
            detail = self.client.get(reverse("app:brand_sources:version_detail", args=[self.source.pk, 1]))
            self.assertEqual(detail.status_code, 200)
            self.assertContains(detail, "Do not write Luma Desk as two words.")
            self.assertEqual(self.client.get(reverse("app:brand_sources:version_detail", args=[self.source.pk, 2])).status_code, 403)
            self.client.logout()
        self.assertEqual(draft.status, SourceVersion.Status.DRAFT)

    def test_draft_mutation_routes_are_server_denied_to_read_only_roles(self):
        rule = self._create_rule()
        fact_text = self.version.source_text + " Up to 20% faster task completion in our internal pilot; results vary."
        self.version.source_text = fact_text
        self.version.save()
        fact = self._create_fact(evidence_excerpt="Up to 20% faster task completion in our internal pilot; results vary.")
        for user in (self.creator, self.reviewer):
            self.client.force_login(user)
            self.assertEqual(self.client.post(reverse("app:brand_sources:rule_delete", args=[self.source.pk, 1, rule.pk])).status_code, 403)
            self.assertEqual(self.client.post(reverse("app:brand_sources:fact_delete", args=[self.source.pk, 1, fact.pk])).status_code, 403)
            self.client.logout()

    def test_known_source_or_rule_ids_cannot_cross_workspace(self):
        other_workspace = Workspace.objects.create(name="Other", slug="brand-source-other")
        outsider = self._member("other_manager", WorkspaceMembership.Role.BRAND_MANAGER, workspace=other_workspace)
        rule = self._create_rule()
        self.client.force_login(outsider)
        self.assertEqual(self.client.get(reverse("app:brand_sources:detail", args=[self.source.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("app:brand_sources:version_edit", args=[self.source.pk, 1])).status_code, 404)
        self.assertEqual(self.client.get(reverse("app:brand_sources:rule_edit", args=[self.source.pk, 1, rule.pk])).status_code, 404)
        response = self.client.post(reverse("app:brand_sources:approve", args=[self.source.pk, 1]))
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.version.status, SourceVersion.Status.DRAFT)

    def test_invalid_source_version_and_nested_entry_ids_return_not_found(self):
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get(reverse("app:brand_sources:detail", args=["00000000-0000-0000-0000-000000000000"])).status_code, 404)
        self.assertEqual(self.client.get(reverse("app:brand_sources:version_detail", args=[self.source.pk, 999])).status_code, 404)
        self.assertEqual(self.client.get(reverse("app:brand_sources:rule_edit", args=[self.source.pk, 1, "00000000-0000-0000-0000-000000000000"])).status_code, 404)

    def test_navigation_destination_resolves_for_every_workspace_role(self):
        expected_path = reverse("app:brand_sources:list")
        for user in (self.creator, self.reviewer, self.manager, self.admin):
            self.client.force_login(user)
            overview = self.client.get(reverse("app:home"))
            self.assertContains(overview, f'href="{expected_path}"')
            page = self.client.get(expected_path)
            self.assertEqual(page.status_code, 200)
            self.client.logout()


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"], DEBUG=True)
class BrandDemoSeedTests(TestCase):
    def setUp(self):
        self.env = {
            "BRANDGUARD_DEMO_PASSWORD": "",
            "BRANDGUARD_DEMO_CREATOR_PASSWORD": "",
            "BRANDGUARD_DEMO_BRAND_MANAGER_PASSWORD": "",
            "BRANDGUARD_DEMO_REVIEWER_PASSWORD": "",
            "BRANDGUARD_DEMO_ADMIN_PASSWORD": "",
        }

    def test_seed_adds_synthetic_sources_without_modifying_them_on_repeat(self):
        from unittest.mock import patch

        with patch.dict("os.environ", self.env):
            call_command("seed_demo", verbosity=0, stdout=io.StringIO())
            workspace = Workspace.objects.get(slug="lumadesk-demo")
            self.assertEqual(workspace.brand_sources.count(), len(LUMADESK_SOURCE_PACK))
            first_revision = workspace.policy_revision
            first_ids = list(workspace.brand_sources.values_list("id", flat=True))
            first_counts = [source.versions.count() for source in workspace.brand_sources.all()]
            self.assertEqual(first_revision, len(LUMADESK_SOURCE_PACK))
            self.assertEqual(workspace.brand_sources.get(slug="lumadesk-product-facts").versions.get().facts.get().identifier, "FACT-02")
            self.assertEqual(workspace.brand_sources.get(slug="lumadesk-product-facts").versions.get().rules.get().rule_type, BrandRule.RuleType.APPROVED_FACT)
            call_command("seed_demo", verbosity=0, stdout=io.StringIO())
        workspace.refresh_from_db()
        self.assertEqual(workspace.policy_revision, first_revision)
        self.assertEqual(list(workspace.brand_sources.values_list("id", flat=True)), first_ids)
        self.assertEqual([source.versions.count() for source in workspace.brand_sources.all()], first_counts)

    def test_seed_refuses_a_non_demo_source_collision(self):
        from unittest.mock import patch

        workspace = Workspace.objects.create(name="LumaDesk", slug="lumadesk-demo", is_demo=True)
        user = get_user_model().objects.create_user(username="lumadesk_brand_manager")
        WorkspaceMembership.objects.create(
            user=user,
            workspace=workspace,
            role=WorkspaceMembership.Role.BRAND_MANAGER,
            is_demo=True,
        )
        BrandSource.objects.create(
            workspace=workspace, owner=user, slug="lumadesk-product-naming", title="Manager source",
            source_type=BrandSource.SourceType.POLICY, is_demo=False,
        )
        with patch.dict("os.environ", self.env), self.assertRaises(CommandError):
            call_command("seed_demo", verbosity=0, stdout=io.StringIO())
        self.assertEqual(workspace.brand_sources.count(), 1)
        self.assertEqual(workspace.brand_sources.get().title, "Manager source")


def _synthetic_pdf(pages):
    """Build a small text PDF fixture without an extra test dependency."""
    encoded_pages = []
    for text in pages:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        encoded_pages.append(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1"))
    font_id = 3 + len(pages) * 2
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    page_ids = [3 + index * 2 for index in range(len(pages))]
    kids = b" ".join(f"{page_id} 0 R".encode("ascii") for page_id in page_ids)
    objects.append(b"<< /Type /Pages /Kids [" + kids + f"] /Count {len(pages)} >>".encode("ascii"))
    for index, content in enumerate(encoded_pages):
        page_id = page_ids[index]
        content_id = page_id + 1
        page = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] ".encode("ascii")
            + f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode("ascii")
        )
        objects.extend([page, b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream"])
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(output)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class BrandDocumentImportWorkflowTests(TestCase):
    source_text = "Official product name: LumaDesk. Internal pilot: up to 20% faster task completion; results vary."

    def setUp(self):
        self.workspace = Workspace.objects.create(name="LumaDesk", slug="document-import-test")
        self.manager = self._member("import_manager", WorkspaceMembership.Role.BRAND_MANAGER)
        self.admin = self._member("import_admin", WorkspaceMembership.Role.ADMIN)
        self.creator = self._member("import_creator", WorkspaceMembership.Role.CREATOR)
        self.reviewer = self._member("import_reviewer", WorkspaceMembership.Role.REVIEWER)

    def tearDown(self):
        storage = storages["private_documents"]
        for name in BrandDocumentImport.objects.values_list("original_document", flat=True):
            if name:
                storage.delete(name)
        super().tearDown()

    def _member(self, username, role, workspace=None):
        user = get_user_model().objects.create_user(
            username=username,
            email=f"{username}@example.test",
            password=secrets.token_urlsafe(24),
        )
        WorkspaceMembership.objects.create(user=user, workspace=workspace or self.workspace, role=role)
        return user

    def _upload(self, name="guide.txt", content=None, content_type="application/octet-stream"):
        return SimpleUploadedFile(name, content if content is not None else self.source_text.encode("utf-8"), content_type)

    def _import(self, *, owner=None, upload=None):
        return import_source_document(
            workspace=self.workspace,
            owner=owner or self.manager,
            title="LumaDesk source",
            source_type=BrandSource.SourceType.GUIDELINES,
            description="Synthetic import fixture.",
            upload=upload or self._upload(),
        )

    def test_txt_import_creates_a_real_draft_and_private_provenance(self):
        raw = b"\xef\xbb\xbf" + self.source_text.encode("utf-8")
        source, version, record = self._import(upload=self._upload("..\\brand-guide.txt", raw, "text/plain"))
        self.assertEqual(source.workspace, self.workspace)
        self.assertEqual(version.status, SourceVersion.Status.DRAFT)
        self.assertEqual(version.source_text, self.source_text)
        self.assertEqual(record.document_type, BrandDocumentImport.DocumentType.TXT)
        self.assertEqual(record.original_filename, "brand-guide.txt")
        self.assertEqual(record.file_sha256, hashlib.sha256(raw).hexdigest())
        self.assertTrue(record.original_document.name.startswith(f"brand-source-imports/{self.workspace.pk}/"))
        self.assertNotIn("brand-guide", record.original_document.name)
        self.assertFalse(record.text_is_confirmed)

    def test_docx_extracts_paragraphs_and_table_text_in_document_order(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Official naming: LumaDesk.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Pilot result"
        table.cell(0, 1).text = "Up to 20% faster; results vary."
        buffer = io.BytesIO()
        document.save(buffer)
        extracted = extract_uploaded_document(self._upload("guide.docx", buffer.getvalue()))
        self.assertEqual(extracted.document_type, "DOCX")
        self.assertEqual(extracted.text, "Official naming: LumaDesk.\nPilot result\tUp to 20% faster; results vary.")

    def test_text_pdf_extracts_page_text_and_preserves_page_boundary(self):
        pdf = _synthetic_pdf(["LumaDesk is the official name.", "Results vary; see the internal pilot."])
        extracted = extract_uploaded_document(self._upload("guide.pdf", pdf))
        self.assertEqual(extracted.document_type, "PDF")
        self.assertEqual(extracted.page_count, 2)
        self.assertIn("LumaDesk is the official name.", extracted.text)
        self.assertIn("\n\n", extracted.text)

    def test_upload_endpoint_creates_draft_and_confirms_manager_or_admin_access(self):
        for user in (self.manager, self.admin):
            self.client.force_login(user)
            response = self.client.get(reverse("app:brand_sources:import"))
            self.assertEqual(response.status_code, 200)
        self.client.force_login(self.manager)
        response = self.client.post(reverse("app:brand_sources:import"), {
            "title": "Uploaded guide", "source_type": BrandSource.SourceType.GUIDELINES,
            "description": "Imported from a synthetic fixture.", "document": self._upload(),
        })
        source = BrandSource.objects.get(title="Uploaded guide")
        version = source.versions.get()
        target = reverse("app:brand_sources:version_edit", args=[source.pk, version.version_number])
        self.assertRedirects(response, target, fetch_redirect_response=False)
        self.assertEqual(version.status, SourceVersion.Status.DRAFT)
        self.assertContains(self.client.get(target), "Document text extracted. Review and confirm the content before publishing.")

    def test_creator_and_reviewer_cannot_import(self):
        for user in (self.creator, self.reviewer):
            self.client.force_login(user)
            self.assertEqual(self.client.get(reverse("app:brand_sources:import")).status_code, 403)
            self.assertEqual(self.client.post(reverse("app:brand_sources:import"), {}).status_code, 403)

    def test_file_errors_are_rejected_without_creating_policy_records(self):
        bad_uploads = [
            self._upload("guide.rtf", b"some text"),
            self._upload("guide.txt", b"\xff\xfe not utf8"),
            self._upload("guide.docx", b"not a zip"),
            self._upload("guide.txt", b"%PDF-1.4 wrong extension"),
            self._upload("guide.pdf", _synthetic_pdf([""])),
            self._upload("guide.pdf", b"%PDF-1.4 malformed"),
        ]
        for upload in bad_uploads:
            with self.subTest(filename=upload.name), self.assertRaises(DocumentImportError):
                self._import(upload=upload)
        self.assertEqual(BrandSource.objects.count(), 0)
        self.assertEqual(BrandDocumentImport.objects.count(), 0)

    def test_upload_size_limit_is_enforced_by_extractor(self):
        with self.assertRaisesMessage(DocumentImportError, "5 MB upload limit"):
            self._import(upload=self._upload("guide.txt", b"x" * (5 * 1024 * 1024 + 1)))

    def test_docx_expansion_limits_and_pdf_page_limit_are_enforced(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "x")
            archive.writestr("word/document.xml", "a" * (10 * 1024 * 1024 + 1))
        with self.assertRaisesMessage(DocumentImportError, "safe processing limit"):
            extract_uploaded_document(self._upload("guide.docx", buffer.getvalue()))
        with self.assertRaisesMessage(DocumentImportError, "limit is 10"):
            extract_uploaded_document(self._upload("guide.pdf", _synthetic_pdf(["Page"] * 11)))

    def test_password_protected_pdf_is_rejected(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.encrypt("synthetic-password")
        buffer = io.BytesIO()
        writer.write(buffer)
        with self.assertRaisesMessage(DocumentImportError, "Password-protected PDFs are not supported"):
            extract_uploaded_document(self._upload("locked.pdf", buffer.getvalue()))

    def test_confirmation_is_required_and_editing_invalidates_it_before_publish(self):
        source, version, record = self._import()
        with self.assertRaisesMessage(ValidationError, "confirm its current version"):
            approve_source_version(source_id=source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        self.client.force_login(self.manager)
        edit_url = reverse("app:brand_sources:version_edit", args=[source.pk, 1])
        response = self.client.post(edit_url, {"source_text": self.source_text + " Corrected.", "effective_date": ""})
        self.assertRedirects(response, edit_url)
        record.refresh_from_db()
        self.assertFalse(record.text_is_confirmed)
        version.refresh_from_db()
        confirm_url = reverse("app:brand_sources:confirm_import", args=[source.pk, 1])
        self.client.post(confirm_url, {"confirm_text": "on"})
        record.refresh_from_db()
        self.assertTrue(record.text_is_confirmed)
        approve_source_version(source_id=source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        version.refresh_from_db()
        self.assertEqual(version.status, SourceVersion.Status.APPROVED)

    def test_confirmation_is_post_only_and_cross_workspace_import_is_not_found(self):
        source, version, _ = self._import()
        self.client.force_login(self.manager)
        confirm_url = reverse("app:brand_sources:confirm_import", args=[source.pk, 1])
        self.assertEqual(self.client.get(confirm_url).status_code, 405)
        other_workspace = Workspace.objects.create(name="Other", slug="other-import-workspace")
        other_manager = self._member("other_import_manager", WorkspaceMembership.Role.BRAND_MANAGER, other_workspace)
        self.client.force_login(other_manager)
        self.assertEqual(self.client.post(confirm_url, {"confirm_text": "on"}).status_code, 404)
        self.assertEqual(BrandSource.objects.count(), 1)

    def test_confirmation_and_approval_publish_the_same_version_to_policy_snapshot(self):
        source, version, record = self._import()
        policy_before = policy_snapshot(self.workspace)
        confirm_imported_text(source_id=source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        approved = approve_source_version(source_id=source.pk, version_number=1, workspace=self.workspace, actor=self.manager)
        self.workspace.refresh_from_db()
        record.refresh_from_db()
        policy_after = policy_snapshot(self.workspace)
        self.assertEqual(approved.pk, version.pk)
        self.assertNotEqual(policy_before["hash"], policy_after["hash"])
        self.assertEqual(policy_after["revision"], policy_before["revision"] + 1)
        self.assertEqual(list(active_policy_versions(self.workspace).values_list("pk", flat=True)), [version.pk])
        self.assertTrue(record.text_is_confirmed)

    def test_imported_rule_and_fact_use_existing_evidence_validation(self):
        source, version, _ = self._import()
        fact = ApprovedFact(
            source_version=version, identifier="FACT-02", subject="Task completion speed",
            statement="Up to 20% faster task completion; results vary.",
            comparison_operator=ApprovedFact.ComparisonOperator.UP_TO, comparison_value=Decimal("20"),
            comparison_unit="%", evidence_excerpt="up to 20% faster task completion; results vary.",
        )
        fact.full_clean()
        fact.save()
        bad_rule = BrandRule(
            source_version=version, identifier="TERM-02", category=BrandRule.Category.TERMINOLOGY,
            severity=BrandRule.Severity.HIGH, description="Use LumaDesk.",
            rule_type=BrandRule.RuleType.PROHIBITED_TERM,
            matcher_data={"type": "PROHIBITED_TERM", "term": "Luma Desk", "case_sensitive": False, "scope": "ALL"},
            evidence_excerpt="This passage does not occur.",
        )
        with self.assertRaises(ValidationError):
            bad_rule.full_clean()
        self.assertEqual(source.workspace_id, version.source.workspace_id)

    def test_malformed_text_no_text_and_extracted_character_limit_are_rejected(self):
        with self.assertRaisesMessage(DocumentImportError, "No meaningful text"):
            extract_uploaded_document(self._upload("empty.txt", b" \n\t"))
        with override_settings(BRAND_DOCUMENT_MAX_EXTRACTED_CHARS=16):
            with self.assertRaisesMessage(DocumentImportError, "character limit"):
                extract_uploaded_document(self._upload("long.txt", self.source_text.encode()))

    def test_database_failure_rolls_back_source_and_removes_private_file(self):
        from unittest.mock import patch

        with patch("brand_sources.services.BrandDocumentImport.save", side_effect=IntegrityError("synthetic failure")):
            with self.assertRaises(DocumentImportPersistenceError):
                self._import()
        self.assertEqual(BrandSource.objects.count(), 0)
        self.assertEqual(BrandDocumentImport.objects.count(), 0)
        storage = storages["private_documents"]
        try:
            _, files = storage.listdir(f"brand-source-imports/{self.workspace.pk}")
        except FileNotFoundError:
            files = []
        self.assertEqual(files, [])

    def test_private_storage_failure_returns_a_safe_error_without_policy_records(self):
        from unittest.mock import patch

        storage = storages["private_documents"]
        with patch.object(storage, "save", side_effect=OSError("synthetic storage failure")):
            with self.assertRaisesMessage(DocumentImportPersistenceError, "could not be stored safely"):
                self._import()
        self.assertEqual(BrandSource.objects.count(), 0)
        self.assertEqual(BrandDocumentImport.objects.count(), 0)
